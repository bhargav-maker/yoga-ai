import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("python-pptx not installed yet")
    sys.exit(1)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
C_DARK = RGBColor(15, 23, 42)        # Slate 900
C_PRIMARY = RGBColor(16, 185, 129)   # Emerald 500
C_ACCENT = RGBColor(14, 165, 233)    # Sky 500
C_WHITE = RGBColor(255, 255, 255)
C_MUTED = RGBColor(148, 163, 184)    # Slate 400
C_CARD_BG = RGBColor(30, 41, 59)     # Slate 800

def set_slide_background(slide, color=C_DARK):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, category="YOGA AI • EXHIBITION 2026"):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_PRIMARY

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = C_WHITE

def create_card(slide, left, top, width, height, title, text_items, accent_color=C_PRIMARY):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD_BG
    card.line.color.rgb = accent_color
    card.line.width = Pt(1.5)

    txBox = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.3), width - Inches(0.6), height - Inches(0.6))
    tf = txBox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = accent_color
    p0.space_after = Pt(14)

    for item in text_items:
        p = tf.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_after = Pt(8)

# ----------------- SLIDE 1: Title Slide -----------------
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide1, C_DARK)

badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(3.2), Inches(0.45))
badge.fill.solid()
badge.fill.fore_color.rgb = C_CARD_BG
badge.line.color.rgb = C_PRIMARY
tf_b = badge.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "AI-POWERED WELLNESS TECH"
p_b.font.size = Pt(11)
p_b.font.bold = True
p_b.font.color.rgb = C_PRIMARY
p_b.alignment = PP_ALIGN.CENTER

t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.0))
tf_t = t_box.text_frame
p_t = tf_t.paragraphs[0]
p_t.text = "Yoga AI: Poddar Yogis"
p_t.font.size = Pt(44)
p_t.font.bold = True
p_t.font.color.rgb = C_WHITE

p_sub = tf_t.add_paragraph()
p_sub.text = "Real-Time Computer Vision & Deep Learning Pose Trainer in the Browser"
p_sub.font.size = Pt(20)
p_sub.font.color.rgb = C_ACCENT
p_sub.space_before = Pt(12)

meta_box = slide1.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.2))
tf_meta = meta_box.text_frame
p_m = tf_meta.paragraphs[0]
p_m.text = "Presented by: Bhargav & Team  |  Exhibition 2026"
p_m.font.size = Pt(15)
p_m.font.color.rgb = C_MUTED
p_m2 = tf_meta.add_paragraph()
p_m2.text = "GitHub: https://github.com/bhargav-maker/yoga-ai"
p_m2.font.size = Pt(13)
p_m2.font.color.rgb = C_PRIMARY

# ----------------- SLIDE 2: Problem & Motivation -----------------
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide2)
add_header(slide2, "The Challenge: Safe & Effective Home Yoga Practice")

create_card(slide2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), 
            "1. Incorrect Form & Injury", 
            ["Performing asanas without supervision often leads to joint strain and chronic back injuries.",
             "Users cannot see their own posture flaws while stretching or looking away from screens.",
             "Lack of posture feedback prevents building proper muscle memory."], 
            RGBColor(239, 68, 68))

create_card(slide2, Inches(4.85), Inches(1.8), Inches(3.6), Inches(4.8), 
            "2. High Cost of Trainers", 
            ["One-on-one personal yoga instructors are expensive and time-inconvenient.",
             "Generic recorded YouTube videos lack interactive correction or validation.",
             "Over 70% of fitness beginners drop out within 3 weeks due to lack of guidance."], 
            RGBColor(245, 158, 11))

create_card(slide2, Inches(8.9), Inches(1.8), Inches(3.6), Inches(4.8), 
            "3. Privacy & Cloud Latency", 
            ["Traditional AI services stream webcam feeds to remote cloud servers, causing privacy risks.",
             "Cloud video latency (300ms+) makes real-time posture correction choppy and frustrating.",
             "Need for a 100% private, on-device local solution."], 
            C_ACCENT)

# ----------------- SLIDE 3: The Solution -----------------
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide3)
add_header(slide3, "Our Solution: Yoga AI (Poddar Yogis)")

create_card(slide3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
            "Real-Time Skeleton Tracking",
            ["Google MoveNet Lightning model tracks 17 skeletal keypoints per frame.",
             "Smooth 30-60 FPS performance without requiring high-end GPUs.",
             "Dynamic canvas aspect-ratio sync aligns landmarks with the user's video feed."],
            C_PRIMARY)

create_card(slide3, Inches(6.9), Inches(1.8), Inches(5.6), Inches(2.3),
            "100% Client-Side Private AI",
            ["Inference runs entirely in the browser using TensorFlow.js.",
             "Webcam video NEVER leaves the user's computer — maximum privacy.",
             "Zero cloud latency and works completely offline."],
            C_ACCENT)

create_card(slide3, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.3),
            "Custom Pose Classification",
            ["Custom deep neural network evaluates normalized 34-dimensional feature vector.",
             "Instant classification of 7 classical yoga asanas.",
             "High-accuracy thresholding prevents false-positive detections."],
            RGBColor(168, 85, 247))

create_card(slide3, Inches(6.9), Inches(4.4), Inches(5.6), Inches(2.3),
            "Gamified Hold Timer & Audio Feedback",
            ["Tracks hold streak duration in seconds when correct posture is maintained.",
             "Audible countdown signals success and motivates continuous practice.",
             "Saves personal best hold times across training sessions."],
            RGBColor(236, 72, 153))

# ----------------- SLIDE 4: Architecture -----------------
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide4)
add_header(slide4, "System Architecture: End-to-End Pipeline")

steps = [
    ("1. Video Capture", "React-Webcam captures standard 640x480 video stream directly from device camera.", C_PRIMARY),
    ("2. MoveNet Lightning", "Extracts 17 keypoints (nose, shoulders, elbows, hips, knees, ankles) in real time.", C_ACCENT),
    ("3. Normalization", "Calculates bounding box and normalizes coordinates, making detection distance-invariant.", RGBColor(168, 85, 247)),
    ("4. Classifier DNN", "Multi-layer dense neural network computes probability distribution across yoga asanas.", RGBColor(245, 158, 11)),
    ("5. Live Feedback", "Draws skeletal overlay on HTML5 canvas and updates hold counter & audio cues.", RGBColor(16, 185, 129))
]

for idx, (stitle, sdesc, scolor) in enumerate(steps):
    top_pos = Inches(1.8 + idx * 1.0)
    card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(0.85))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD_BG
    card.line.color.rgb = scolor
    card.line.width = Pt(1.5)

    tx = slide4.shapes.add_textbox(Inches(1.1), top_pos + Inches(0.12), Inches(11.1), Inches(0.6))
    tf = tx.text_frame
    p = tf.paragraphs[0]
    p.text = stitle + "  —  "
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = scolor

    r = p.add_run()
    r.text = sdesc
    r.font.bold = False
    r.font.color.rgb = C_WHITE

# ----------------- SLIDE 5: Supported Poses -----------------
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide5)
add_header(slide5, "7 Supported Classical Asanas")

poses = [
    ("Tree Pose (Vrikshasana)", "Improves balance, posture, and core leg stability."),
    ("Chair Pose (Utkatasana)", "Strengthens thighs, spine, and activates calves."),
    ("Cobra Pose (Bhujangasana)", "Enhances spinal flexibility and chest expansion."),
    ("Warrior Pose (Virabhadrasana)", "Develops stamina, focus, and hip mobility."),
    ("Downward Dog (Adho Mukha)", "Full body stretch, decompresses spine & shoulders."),
    ("Shoulderstand (Sarvangasana)", "Boosts circulation, thyroid function & upper body."),
    ("Triangle Pose (Trikonasana)", "Stretches hamstrings, groin, and stimulates organs.")
]

for i, (pname, pdesc) in enumerate(poses):
    col = 0 if i < 4 else 1
    row = i if i < 4 else (i - 4)
    l = Inches(0.8) if col == 0 else Inches(6.9)
    t = Inches(1.8 + row * 1.25)
    w = Inches(5.6)
    h = Inches(1.1)

    c = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    c.fill.solid()
    c.fill.fore_color.rgb = C_CARD_BG
    c.line.color.rgb = C_PRIMARY
    c.line.width = Pt(1)

    tb = slide5.shapes.add_textbox(l + Inches(0.25), t + Inches(0.15), w - Inches(0.5), h - Inches(0.3))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "🧘 " + pname
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p2 = tf.add_paragraph()
    p2.text = pdesc
    p2.font.size = Pt(12)
    p2.font.color.rgb = C_MUTED

# ----------------- SLIDE 6: Tech Stack & Engineering -----------------
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide6)
add_header(slide6, "Technical Stack & Engineering Highlights")

create_card(slide6, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
            "Frontend & UI",
            ["React 18 & React Router for responsive single-page architecture.",
             "HTML5 Canvas rendering engine for low-latency skeletal overlay.",
             "Modern CSS Glassmorphism styling with dark theme.",
             "Custom Audio countdown cues via Web Audio API."],
            C_PRIMARY)

create_card(slide6, Inches(4.85), Inches(1.8), Inches(3.6), Inches(4.8),
            "AI & Computer Vision",
            ["TensorFlow.js MoveNet Lightning: real-time pose detector.",
             "Custom Keras/TensorFlow DNN Classifier.",
             "Offline Bundled Model Weights (model.json + shard binaries).",
             "Joint score filtering (>0.3) for clean, jitter-free lines."],
            C_ACCENT)

create_card(slide6, Inches(8.9), Inches(1.8), Inches(3.6), Inches(4.8),
            "Deployment & Tools",
            ["One-Click Desktop Launcher (start.bat / yoga.bat).",
             "Instant Production Server (npx serve / Node CLI).",
             "Open Source on GitHub: bhargav-maker/yoga-ai.",
             "Cross-Platform: Windows, Mac, Linux, and Mobile."],
            RGBColor(168, 85, 247))

# ----------------- SLIDE 7: Future Roadmap -----------------
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide7)
add_header(slide7, "Future Roadmap: Next Generation Yoga AI")

create_card(slide7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
            "1. Real-Time Angle Correction",
            ["Mathematical angle computation (e.g. knee 90°, elbow 180°).",
             "Visual color-coded guidance (Red = bad angle, Green = perfect).",
             "Hyper-targeted feedback on which specific limb to adjust."],
            C_PRIMARY)

create_card(slide7, Inches(6.9), Inches(1.8), Inches(5.6), Inches(2.3),
            "2. Voice AI Coach",
            ["Integrated text-to-speech AI instructor giving spoken corrections.",
             "Hands-free guidance: 'Lift your right leg higher', 'Straighten back'.",
             "Zero need to glance at screen during balancing postures."],
            C_ACCENT)

create_card(slide7, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.3),
            "3. Multi-Person & Group Sessions",
            ["MoveNet MultiPose integration for classroom and family workouts.",
             "Leaderboards, multiplayer pose battles, and social challenges.",
             "Class analytics dashboard for yoga instructors."],
            RGBColor(245, 158, 11))

create_card(slide7, Inches(6.9), Inches(4.4), Inches(5.6), Inches(2.3),
            "4. Mobile App & Wearables",
            ["Flutter / React Native cross-platform mobile apps.",
             "Smartwatch heart rate and calorie synchronization.",
             "Personalized daily wellness recommendations based on fatigue."],
            RGBColor(236, 72, 153))

# ----------------- SLIDE 8: Conclusion & Thank You -----------------
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide8)

tb_c = slide8.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(2.0))
tfc = tb_c.text_frame
p = tfc.paragraphs[0]
p.text = "Thank You!"
p.font.size = Pt(46)
p.font.bold = True
p.font.color.rgb = C_PRIMARY

p2 = tfc.add_paragraph()
p2.text = "Yoga AI: Empowering Personal Wellness with Private On-Device AI"
p2.font.size = Pt(22)
p2.font.color.rgb = C_WHITE
p2.space_before = Pt(12)

box_demo = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.8), Inches(11.7), Inches(2.6))
box_demo.fill.solid()
box_demo.fill.fore_color.rgb = C_CARD_BG
box_demo.line.color.rgb = C_ACCENT
box_demo.line.width = Pt(1.5)

tfd = box_demo.text_frame
pd1 = tfd.paragraphs[0]
pd1.text = "Ready for Live Demonstration!"
pd1.font.size = Pt(22)
pd1.font.bold = True
pd1.font.color.rgb = C_ACCENT
pd1.space_after = Pt(14)

pd2 = tfd.add_paragraph()
pd2.text = "• Live Web App: http://localhost:3000"
pd2.font.size = Pt(16)
pd2.font.color.rgb = C_WHITE
pd2.space_after = Pt(8)

pd3 = tfd.add_paragraph()
pd3.text = "• Open Source Repository: https://github.com/bhargav-maker/yoga-ai"
pd3.font.size = Pt(16)
pd3.font.color.rgb = C_PRIMARY
pd3.space_after = Pt(8)

pd4 = tfd.add_paragraph()
pd4.text = "• Questions & Feedback Welcomed!"
pd4.font.size = Pt(16)
pd4.font.color.rgb = C_MUTED

output_path = os.path.join(os.getcwd(), "Yoga_AI_Presentation.pptx")
prs.save(output_path)
print(f"SUCCESS: Saved presentation to {output_path}")
